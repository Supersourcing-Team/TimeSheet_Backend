import io
import csv
from datetime import date, datetime
from typing import Any, Dict, List, Optional
from sqlalchemy import select
from sqlalchemy.ext.asyncio import AsyncSession
from sqlalchemy.orm import selectinload
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from app.modules.timesheets.model import Timesheet
from app.modules.users.model import User
from app.modules.departments.model import Department
from app.modules.project_assignments.model import ProjectAssignment
from app.modules.projects.model import Project
from app.modules.roles.model import Role
from app.modules.leave_types.model import LeaveType
from app.modules.leave_balances.model import LeaveBalance
from app.modules.leave_requests.model import LeaveRequest
from app.modules.milestones.model import Milestone
from app.modules.milestones.assignment_model import MilestoneAssignment
from app.modules.weekend_work.model import WeekendWorkRequest
from app.modules.clients.model import Client
from app.modules.tool_allocations.model import ToolAllocation
from app.modules.tools.model import Tool


class EodReportService:

    @staticmethod
    async def get_daily_eod_entries(db: AsyncSession, target_date: date) -> List[Dict[str, Any]]:
        """
        Fetches all submitted timesheet entries for a specific date with joined
        employee, department, project, and assignment details.
        """
        stmt = (
            select(Timesheet)
            .options(
                selectinload(Timesheet.user).selectinload(User.department),
                selectinload(Timesheet.project_assignment).selectinload(ProjectAssignment.project),
            )
            .where(Timesheet.timesheet_date == target_date)
            .order_by(Timesheet.user_id, Timesheet.id)
        )
        result = await db.execute(stmt)
        timesheets = result.scalars().all()

        entries: List[Dict[str, Any]] = []
        for t in timesheets:
            user = t.user
            dept_name = user.department.name if (user and user.department) else "General"
            proj = t.project_assignment.project if (t.project_assignment and t.project_assignment.project) else None
            proj_name = proj.project_name if proj else "Unknown Project"
            proj_code = f"PRJ-{proj.id}" if proj else ""

            b_hours = float(t.billable_hours or 0.0)
            nb_hours = float(t.non_billable_hours or 0.0)
            tot_hours = round(b_hours + nb_hours, 2)

            entries.append({
                "timesheet_id": t.id,
                "date": str(t.timesheet_date),
                "employee_id": t.user_id,
                "employee_name": f"{user.first_name} {user.last_name}" if user else "Unknown",
                "employee_email": user.email if user else "",
                "department": dept_name,
                "project_id": proj.id if proj else None,
                "project_name": proj_name,
                "project_code": proj_code,
                "billable_hours": b_hours,
                "billable_work_summary": t.billable_work_summary or "",
                "non_billable_hours": nb_hours,
                "non_billable_work_summary": t.non_billable_work_summary or "",
                "total_hours": tot_hours,
                "status": t.status,
                "submitted_at": t.created_at.strftime("%Y-%m-%d %H:%M:%S") if t.created_at else "",
            })

        return entries

    @staticmethod
    def generate_daily_eod_excel(entries: List[Dict[str, Any]], target_date: date) -> bytes:
        """
        Generates a formatted Excel (.xlsx) file containing all employees' separate
        submitted EOD work descriptions and hours for a single day.
        """
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = f"EOD_{target_date.strftime('%Y%m%d')}"

        # Setup styles
        title_font = Font(name="Calibri", size=16, bold=True, color="1E293B")
        subtitle_font = Font(name="Calibri", size=11, italic=True, color="64748B")
        header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
        data_font = Font(name="Calibri", size=10, color="0F172A")
        bold_font = Font(name="Calibri", size=10, bold=True, color="0F172A")
        total_font = Font(name="Calibri", size=11, bold=True, color="1E293B")

        header_fill = PatternFill(start_color="1E293B", end_color="1E293B", fill_type="solid")
        billable_header_fill = PatternFill(start_color="065F46", end_color="065F46", fill_type="solid")
        non_billable_header_fill = PatternFill(start_color="334155", end_color="334155", fill_type="solid")
        
        zebra_fill = PatternFill(start_color="F8FAFC", end_color="F8FAFC", fill_type="solid")
        white_fill = PatternFill(start_color="FFFFFF", end_color="FFFFFF", fill_type="solid")
        total_fill = PatternFill(start_color="E2E8F0", end_color="E2E8F0", fill_type="solid")

        thin_side = Side(style="thin", color="CBD5E1")
        border = Border(left=thin_side, right=thin_side, top=thin_side, bottom=thin_side)
        double_bottom_side = Side(style="double", color="1E293B")
        total_border = Border(top=thin_side, bottom=double_bottom_side, left=thin_side, right=thin_side)

        align_center = Alignment(horizontal="center", vertical="center")
        align_left = Alignment(horizontal="left", vertical="center", wrap_text=True)
        align_right = Alignment(horizontal="right", vertical="center")
        align_header = Alignment(horizontal="center", vertical="center", wrap_text=True)

        # Title block
        ws.merge_cells("A1:L1")
        ws["A1"] = f"DAILY EOD TIMESHEET REPORT — {target_date.strftime('%A, %B %d, %Y')}"
        ws["A1"].font = title_font
        ws["A1"].alignment = Alignment(horizontal="left", vertical="center")

        ws.merge_cells("A2:L2")
        total_employees = len(set(e["employee_id"] for e in entries))
        total_b = sum(e["billable_hours"] for e in entries)
        total_nb = sum(e["non_billable_hours"] for e in entries)
        grand_tot = sum(e["total_hours"] for e in entries)
        ws["A2"] = (
            f"Generated on: {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}  |  "
            f"Total Submissions: {len(entries)}  |  "
            f"Active Employees: {total_employees}  |  "
            f"Total Billable: {total_b:.2f}h  |  "
            f"Total Non-Billable: {total_nb:.2f}h  |  "
            f"Grand Total: {grand_tot:.2f}h"
        )
        ws["A2"].font = subtitle_font
        ws["A2"].alignment = Alignment(horizontal="left", vertical="center")

        ws.row_dimensions[1].height = 28
        ws.row_dimensions[2].height = 20
        ws.row_dimensions[4].height = 28

        # Header definitions
        headers = [
            ("#", 6, header_fill, align_center),
            ("Employee Name", 22, header_fill, align_left),
            ("Email", 26, header_fill, align_left),
            ("Department", 16, header_fill, align_center),
            ("Project", 22, header_fill, align_left),
            ("Project Code", 14, header_fill, align_center),
            ("Billable (h)", 14, billable_header_fill, align_right),
            ("Billable Work Description / Client Scope", 42, billable_header_fill, align_left),
            ("Non-Billable (h)", 16, non_billable_header_fill, align_right),
            ("Non-Billable Description / Internal Notes", 42, non_billable_header_fill, align_left),
            ("Total (h)", 12, header_fill, align_right),
            ("Status", 14, header_fill, align_center),
        ]

        # Write header
        row_num = 4
        for col_idx, (col_name, col_width, fill_color, _) in enumerate(headers, start=1):
            cell = ws.cell(row=row_num, column=col_idx, value=col_name)
            cell.font = header_font
            cell.fill = fill_color
            cell.alignment = align_header
            cell.border = border
            col_letter = get_column_letter(col_idx)
            ws.column_dimensions[col_letter].width = col_width

        # Write data rows
        row_num = 5
        for idx, item in enumerate(entries, start=1):
            fill = zebra_fill if idx % 2 == 0 else white_fill
            ws.row_dimensions[row_num].height = 36

            row_values = [
                (idx, align_center, data_font),
                (item["employee_name"], align_left, bold_font),
                (item["employee_email"], align_left, data_font),
                (item["department"], align_center, data_font),
                (item["project_name"], align_left, bold_font),
                (item["project_code"], align_center, data_font),
                (f"{item['billable_hours']:.2f}", align_right, bold_font),
                (item["billable_work_summary"], align_left, data_font),
                (f"{item['non_billable_hours']:.2f}", align_right, bold_font),
                (item["non_billable_work_summary"], align_left, data_font),
                (f"{item['total_hours']:.2f}", align_right, bold_font),
                (item["status"].capitalize() if item["status"] else "Submitted", align_center, data_font),
            ]

            for col_idx, (val, alignment, font) in enumerate(row_values, start=1):
                cell = ws.cell(row=row_num, column=col_idx, value=val)
                cell.font = font
                cell.fill = fill
                cell.alignment = alignment
                cell.border = border

            row_num += 1

        # Summary total row
        if entries:
            ws.row_dimensions[row_num].height = 24
            ws.merge_cells(start_row=row_num, start_column=1, end_row=row_num, end_column=6)
            total_label_cell = ws.cell(row=row_num, column=1, value="TOTALS ACROSS ALL EMPLOYEES")
            total_label_cell.font = total_font
            total_label_cell.alignment = Alignment(horizontal="right", vertical="center")
            total_label_cell.fill = total_fill
            
            for col_idx in range(1, 7):
                ws.cell(row=row_num, column=col_idx).fill = total_fill
                ws.cell(row=row_num, column=col_idx).border = total_border

            # Total billable
            b_cell = ws.cell(row=row_num, column=7, value=f"{total_b:.2f}")
            b_cell.font = total_font
            b_cell.alignment = align_right
            b_cell.fill = total_fill
            b_cell.border = total_border

            # Blank desc
            ws.cell(row=row_num, column=8, value="").fill = total_fill
            ws.cell(row=row_num, column=8).border = total_border

            # Total non-billable
            nb_cell = ws.cell(row=row_num, column=9, value=f"{total_nb:.2f}")
            nb_cell.font = total_font
            nb_cell.alignment = align_right
            nb_cell.fill = total_fill
            nb_cell.border = total_border

            # Blank desc
            ws.cell(row=row_num, column=10, value="").fill = total_fill
            ws.cell(row=row_num, column=10).border = total_border

            # Grand total
            tot_cell = ws.cell(row=row_num, column=11, value=f"{grand_tot:.2f}")
            tot_cell.font = total_font
            tot_cell.alignment = align_right
            tot_cell.fill = total_fill
            tot_cell.border = total_border

            # Status cell
            ws.cell(row=row_num, column=12, value="").fill = total_fill
            ws.cell(row=row_num, column=12).border = total_border

        output = io.BytesIO()
        wb.save(output)
        return output.getvalue()

    @staticmethod
    def generate_daily_eod_csv(entries: List[Dict[str, Any]], target_date: date) -> str:
        """
        Generates a standard CSV representation of the daily EOD report.
        """
        output = io.StringIO()
        fieldnames = [
            "Index",
            "Date",
            "Employee Name",
            "Email",
            "Department",
            "Project Name",
            "Project Code",
            "Billable Hours",
            "Billable Work Summary",
            "Non-Billable Hours",
            "Non-Billable Work Summary",
            "Total Hours",
            "Status",
            "Submitted At",
        ]
        writer = csv.writer(output, lineterminator="\n")
        writer.writerow(fieldnames)

        for idx, item in enumerate(entries, start=1):
            writer.writerow([
                idx,
                item["date"],
                item["employee_name"],
                item["employee_email"],
                item["department"],
                item["project_name"],
                item["project_code"],
                f"{item['billable_hours']:.2f}",
                item["billable_work_summary"],
                f"{item['non_billable_hours']:.2f}",
                item["non_billable_work_summary"],
                f"{item['total_hours']:.2f}",
                item["status"],
                item["submitted_at"],
            ])

        return output.getvalue()

    @staticmethod
    async def send_daily_eod_to_slack(db: AsyncSession, target_date: date) -> Dict[str, Any]:
        """
        Gathers all employee EOD submissions for target_date, formats the Slack payload,
        generates the Excel attachment, and dispatches to Slack via SlackService.
        """
        from app.services.slack_service import SlackService
        import logging
        logger = logging.getLogger(__name__)

        entries = await EodReportService.get_daily_eod_entries(db, target_date)
        total_employees = len(set(e["employee_id"] for e in entries))
        total_b = sum(e["billable_hours"] for e in entries)
        total_nb = sum(e["non_billable_hours"] for e in entries)
        grand_tot = sum(e["total_hours"] for e in entries)

        date_str = target_date.strftime("%A, %B %d, %Y")
        short_date = target_date.strftime("%Y%m%d")

        # 1. Clean minimal header comment for file attachment
        file_comment = f"📄 *Daily EOD Timesheet Report — {date_str}*"

        # 2. Generate Excel workbook
        excel_bytes = EodReportService.generate_daily_eod_excel(entries, target_date)
        filename = f"Daily_EOD_Report_{short_date}.xlsx"
        title = f"Daily EOD Report - {date_str}"

        # 3. Upload file directly to Slack channel
        upload_result = await SlackService.upload_file(
            file_bytes=excel_bytes,
            filename=filename,
            title=title,
            initial_comment=file_comment,
        )

        if upload_result.get("ok"):
            return {
                "ok": True,
                "message": f"Daily EOD report file for {date_str} sent successfully to Slack.",
                "details": upload_result,
            }

        # 4. Fallback if file upload fails
        logger.warning(f"File upload to Slack failed ({upload_result.get('error')}), falling back to chat.postMessage...")
        msg_result = await SlackService.post_message(text=f"Daily EOD Timesheet Report for {date_str}")
        if msg_result.get("ok"):
            return {
                "ok": True,
                "message": f"Daily EOD report notification posted to Slack.",
                "details": msg_result,
            }

        return {
            "ok": False,
            "error": upload_result.get("error") or msg_result.get("error") or "Failed to dispatch to Slack",
        }

